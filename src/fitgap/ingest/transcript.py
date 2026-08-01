"""Extract *implied* requirements from workshop transcripts (.vtt / .txt / .docx).

This is the least reliable source, and the pipeline treats it that way:

* Deterministic code parses the transcript into (timestamp, speaker, text) cues.
* The LLM extracts implied requirements per chunk, and must tag each with the
  timestamp/speaker of the utterance that implies it — full traceability.
* Every transcript-derived requirement is marked
  ``source_reliability: inferred`` so it is flagged in the register.
* Cue text is anonymised before it is sent to the API.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from fitgap.classify.classifier import ClassificationError
from fitgap.config import Config
from fitgap.models import (
    FunctionalArea,
    ParsedRequirement,
    RedactionEvent,
    SourceReliability,
    SourceType,
)
from fitgap.redact import RedactionRules, anonymise

CHUNK_CHAR_LIMIT = 6000

VTT_TIMING_RE = re.compile(
    r"^(?P<start>\d{1,2}:\d{2}(?::\d{2})?(?:\.\d{1,3})?)\s*-->\s*"
    r"\d{1,2}:\d{2}(?::\d{2})?(?:\.\d{1,3})?"
)
VOICE_TAG_RE = re.compile(r"<v\s+(?P<speaker>[^>]+)>(?P<text>.*?)(?:</v>)?$", re.DOTALL)
SPEAKER_LINE_RE = re.compile(r"^\[?(?P<ts>\d{1,2}:\d{2}(?::\d{2})?)\]?\s*(?P<rest>.*)$")
SPEAKER_PREFIX_RE = re.compile(r"^(?P<speaker>[A-Za-z][\w .'-]{0,40}):\s+(?P<text>.+)$")

EXTRACT_SYSTEM_PROMPT = """\
You are a business analyst reviewing a Dynamics 365 / Power Platform discovery \
workshop transcript. Extract the IMPLIED BUSINESS REQUIREMENTS from the \
conversation.

Rules:
- A requirement is a capability the client's business needs from the solution. \
Rephrase it as a single standalone declarative statement ("The system must ..."), \
understandable without the transcript.
- For each requirement, record the timestamp and speaker of the utterance that \
implies it, exactly as given in the transcript.
- Be conservative: extract only what the conversation genuinely implies. Do NOT \
invent requirements, do NOT extract pleasantries, logistics, or discussion \
about the workshop itself.
- If the same need is expressed multiple times, extract it once (first mention).
- Speaker names may be redaction placeholders like [PERSON]; keep them as-is.
- Assign the functional_area that best matches each requirement.
"""

EXTRACT_TOOL = {
    "name": "record_extracted_requirements",
    "description": "Record business requirements implied by the transcript chunk.",
    "input_schema": {
        "type": "object",
        "properties": {
            "requirements": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "text": {"type": "string"},
                        "timestamp": {"type": "string"},
                        "speaker": {"type": "string"},
                        "functional_area": {
                            "type": "string",
                            "enum": [a.value for a in FunctionalArea],
                        },
                    },
                    "required": ["text", "timestamp", "speaker"],
                },
            }
        },
        "required": ["requirements"],
    },
}


@dataclass
class Cue:
    timestamp: str  # "hh:mm:ss" ("" when the transcript has no timestamps)
    speaker: str    # "" when unknown
    text: str


def _clean_timestamp(raw: str) -> str:
    return raw.split(".")[0]


def parse_vtt(path: Path) -> list[Cue]:
    cues: list[Cue] = []
    current_ts: str | None = None
    in_note = False
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        stripped = line.strip()
        if not stripped:
            current_ts = None
            in_note = False
            continue
        if stripped == "WEBVTT" or stripped.startswith(("NOTE", "STYLE", "REGION")):
            in_note = True
            continue
        timing = VTT_TIMING_RE.match(stripped)
        if timing:
            current_ts = _clean_timestamp(timing.group("start"))
            in_note = False
            continue
        if in_note or current_ts is None:
            continue  # cue identifiers before a timing line are also skipped
        voice = VOICE_TAG_RE.match(stripped)
        if voice:
            cues.append(
                Cue(current_ts, voice.group("speaker").strip(), voice.group("text").strip())
            )
            continue
        prefixed = SPEAKER_PREFIX_RE.match(stripped)
        if prefixed:
            cues.append(
                Cue(current_ts, prefixed.group("speaker").strip(), prefixed.group("text").strip())
            )
        elif cues and cues[-1].timestamp == current_ts:
            cues[-1].text += " " + stripped  # continuation line of the same cue
        else:
            cues.append(Cue(current_ts, "", stripped))
    return [c for c in cues if c.text]


def parse_txt_lines(lines: list[str]) -> list[Cue]:
    """Parse plain-text transcript lines: '[00:12:34] Speaker: text' variants."""
    cues: list[Cue] = []
    for raw in lines:
        stripped = raw.strip()
        if not stripped:
            continue
        timestamp = ""
        ts_match = SPEAKER_LINE_RE.match(stripped)
        if ts_match and ts_match.group("ts"):
            timestamp = _clean_timestamp(ts_match.group("ts"))
            stripped = ts_match.group("rest").strip()
        speaker = ""
        prefixed = SPEAKER_PREFIX_RE.match(stripped)
        if prefixed:
            speaker = prefixed.group("speaker").strip()
            stripped = prefixed.group("text").strip()
        if stripped:
            cues.append(Cue(timestamp, speaker, stripped))
    return cues


def parse_transcript(path: Path) -> list[Cue]:
    suffix = path.suffix.lower()
    if suffix == ".vtt":
        return parse_vtt(path)
    if suffix == ".txt":
        return parse_txt_lines(path.read_text(encoding="utf-8-sig").splitlines())
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        return parse_txt_lines([p.text for p in doc.paragraphs])
    raise ValueError(f"Unsupported transcript type: {path.name}")


def _chunk(cues: list[Cue], char_limit: int = CHUNK_CHAR_LIMIT) -> list[list[Cue]]:
    chunks: list[list[Cue]] = []
    current: list[Cue] = []
    size = 0
    for cue in cues:
        cue_len = len(cue.text) + len(cue.speaker) + 16
        if current and size + cue_len > char_limit:
            chunks.append(current)
            current, size = [], 0
        current.append(cue)
        size += cue_len
    if current:
        chunks.append(current)
    return chunks


class TranscriptExtractor:
    def __init__(self, config: Config, rules: RedactionRules, client) -> None:
        self.config = config
        self.rules = rules
        self.client = client

    def extract(
        self,
        path: Path,
        on_progress=None,  # callable(done, total) for progress display
    ) -> tuple[list[ParsedRequirement], list[RedactionEvent]]:
        cues = parse_transcript(path)
        if not cues:
            return [], []
        events: list[RedactionEvent] = []
        results: list[ParsedRequirement] = []
        seen_texts: set[str] = set()
        chunks = _chunk(cues)
        if on_progress:
            on_progress(0, len(chunks))
        for chunk_index, chunk in enumerate(chunks, start=1):
            lines = []
            for cue in chunk:
                redacted_speaker, speaker_events = anonymise(cue.speaker, self.rules)
                redacted_text, text_events = anonymise(cue.text, self.rules)
                events.extend(speaker_events)
                events.extend(text_events)
                prefix = f"[{cue.timestamp}] " if cue.timestamp else ""
                who = f"{redacted_speaker}: " if redacted_speaker else ""
                lines.append(f"{prefix}{who}{redacted_text}")
            transcript_block = "\n".join(lines)

            response = self.client.messages.create(
                model=self.config.model,
                max_tokens=4096,
                system=EXTRACT_SYSTEM_PROMPT,
                messages=[
                    {
                        "role": "user",
                        "content": (
                            "Extract the implied business requirements from this "
                            "transcript chunk:\n\n" + transcript_block
                        ),
                    }
                ],
                tools=[EXTRACT_TOOL],
                tool_choice={"type": "tool", "name": "record_extracted_requirements"},
            )
            tool_input = None
            for block in response.content:
                if getattr(block, "type", None) == "tool_use":
                    tool_input = block.input
                    break
            if tool_input is None:
                raise ClassificationError(
                    "Model response contained no record_extracted_requirements "
                    "tool call."
                )
            for entry in tool_input.get("requirements", []):
                text = (entry.get("text") or "").strip()
                if len(text) < 10 or text.lower() in seen_texts:
                    continue
                seen_texts.add(text.lower())
                timestamp = (entry.get("timestamp") or "").strip()
                speaker = (entry.get("speaker") or "").strip()
                where = timestamp or "no-timestamp"
                who = f" ({speaker})" if speaker else ""
                area = entry.get("functional_area")
                results.append(
                    ParsedRequirement(
                        text=text,
                        source=SourceType.TRANSCRIPT,
                        source_ref=f"{path.name}#{where}{who}",
                        functional_area=FunctionalArea(area)
                        if area
                        else FunctionalArea.OTHER,
                        source_reliability=SourceReliability.INFERRED,
                    )
                )
            if on_progress:
                on_progress(chunk_index, len(chunks))
        return results, events
